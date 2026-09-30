"""Lectura del archivo de Excel. Extracción pura: no normaliza ni decide nada."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from scripts.migracion import catalogos as cat

#: Nombres de hoja esperados y su número de fila donde empieza el encabezado (1-indexada).
FILA_ENCABEZADO = 1
HOJA_CHOFERES = "Choferes_y_Moviles"
HOJA_CLIENTES = "Clientes_Cotizaciones"
HOJA_VIAJES = "Registro_Viajes"

HOJAS_ESPERADAS = (HOJA_CHOFERES, HOJA_CLIENTES, HOJA_VIAJES)


class ArchivoDeMigracionError(Exception):
    """El archivo no se puede usar como fuente de la migración."""


@dataclass(frozen=True)
class Fila:
    """Una fila del Excel con su número original, para poder citarla en el reporte."""

    hoja: str
    numero: int
    datos: dict[str, object]


def _a_fila(hoja: str, numero: int, encabezados: list[str], valores: tuple[object, ...]) -> Fila:
    return Fila(
        hoja=hoja,
        numero=numero,
        datos=dict(zip(encabezados, valores, strict=False)),
    )


def leer_filas(ruta: Path, nombre_hoja: str) -> list[Fila]:
    """Devuelve las filas de datos de una hoja, con la primera fila como encabezado."""
    libro = load_workbook(ruta, data_only=True, read_only=True)
    try:
        if nombre_hoja not in libro.sheetnames:
            raise ArchivoDeMigracionError(
                f"El archivo no tiene la hoja {nombre_hoja!r}. Tiene: {', '.join(libro.sheetnames)}"
            )
        hoja = libro[nombre_hoja]
        filas = list(hoja.iter_rows(values_only=True))
    finally:
        libro.close()

    if not filas:
        raise ArchivoDeMigracionError(f"La hoja {nombre_hoja!r} está vacía")

    encabezados = [str(valor).strip() for valor in filas[FILA_ENCABEZADO - 1]]
    resultado = [
        _a_fila(nombre_hoja, numero, encabezados, valores)
        for numero, valores in enumerate(filas[FILA_ENCABEZADO:], start=FILA_ENCABEZADO + 1)
        if _tiene_datos(valores)
    ]
    if not resultado:
        raise ArchivoDeMigracionError(f"La hoja {nombre_hoja!r} no tiene filas de datos")
    return resultado


def _tiene_datos(valores: tuple[object, ...]) -> bool:
    return any(
        valor is not None and (not isinstance(valor, str) or valor.strip()) for valor in valores
    )


def leer_choferes(ruta: Path) -> list[Fila]:
    return leer_filas(ruta, HOJA_CHOFERES)


def leer_clientes(ruta: Path) -> list[Fila]:
    return leer_filas(ruta, HOJA_CLIENTES)


def leer_viajes(ruta: Path) -> list[Fila]:
    return leer_filas(ruta, HOJA_VIAJES)


def verificar_hojas(ruta: Path) -> None:
    """Falla temprano si falta alguna hoja esperada, en vez de migrar la mitad."""
    libro = load_workbook(ruta, read_only=True)
    try:
        faltantes = [hoja for hoja in HOJAS_ESPERADAS if hoja not in libro.sheetnames]
    finally:
        libro.close()
    if faltantes:
        raise ArchivoDeMigracionError(
            f"Faltan hojas en el archivo: {', '.join(faltantes)}. "
            f"Esperadas: {', '.join(HOJAS_ESPERADAS)}"
        )


def ruta_por_defecto() -> Path:
    """El archivo tal como quedó en el repo, en `raw-data/`."""
    raiz = Path(__file__).resolve().parents[3]
    return raiz / "raw-data" / "Caso 2 Fletes_Mudanzas_Express_Datos.xlsx"


def resumen_de_catalogos() -> list[str]:
    """Los catálogos provisionales, para que el reporte los muestre y no queden invisibles."""
    return [
        "Tipos de vehículo: " + ", ".join(cat.TIPOS_VEHICULO_CANONICOS),
        "Estados de móvil: " + ", ".join(cat.ESTADOS_CANONICOS),
        "Condiciones de IVA: " + ", ".join(cat.CONDICIONES_IVA_CANONICAS),
        "Medios de pago (PROPUESTOS, no confirmados): " + ", ".join(cat.MEDIOS_PAGO_PROPUESTOS),
        "Estados de cobro: " + ", ".join(cat.ESTADOS_COBRO_CANONICOS),
    ]
