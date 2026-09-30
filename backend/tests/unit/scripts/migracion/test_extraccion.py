"""Tests de la lectura del archivo de Excel (FR-001, FR-002).

A diferencia del resto de la historia, estos tests sí necesitan el archivo real: no tiene sentido
probar la extracción con un doble, porque lo que se quiere verificar es que las tres hojas del
archivo del cliente se leen bien y que las filas vacías no se cuelan como datos.
"""

from pathlib import Path

import pytest

from scripts.migracion.extraccion import (
    HOJA_CHOFERES,
    HOJA_CLIENTES,
    HOJA_VIAJES,
    ArchivoDeMigracionError,
    leer_choferes,
    leer_clientes,
    leer_filas,
    leer_viajes,
    ruta_por_defecto,
    verificar_hojas,
)

ESPERADAS = {HOJA_CHOFERES: 10, HOJA_CLIENTES: 10, HOJA_VIAJES: 10}


@pytest.fixture(scope="module")
def archivo() -> Path:
    ruta = ruta_por_defecto()

    if not ruta.exists():
        pytest.skip(f"No está el archivo de Excel en {ruta}")

    return ruta


class TestRutaPorDefecto:
    def test_apunta_al_archivo_del_cliente(self) -> None:
        assert ruta_por_defecto().name.endswith(".xlsx")
        assert ruta_por_defecto().parent.name == "raw-data"


class TestVerificarHojas:
    def test_el_archivo_real_tiene_las_tres_hojas(self, archivo: Path) -> None:
        verificar_hojas(archivo)

    def test_una_hoja_que_falta_falla_con_un_mensaje_claro(self, tmp_path: Path) -> None:
        """FR-001: falla temprano y sin escribir nada, en vez de migrar la mitad."""
        from openpyxl import Workbook

        libro = Workbook()
        libro.active.title = HOJA_CLIENTES
        ruta = tmp_path / "incompleto.xlsx"
        libro.save(ruta)

        with pytest.raises(ArchivoDeMigracionError, match="Faltan hojas"):
            verificar_hojas(ruta)

    def test_una_hoja_inexistente_falla_al_leerla(self, archivo: Path) -> None:
        with pytest.raises(ArchivoDeMigracionError, match="no tiene la hoja"):
            leer_filas(archivo, "Hoja Que No Existe")


class TestLectura:
    @pytest.mark.parametrize(
        ("leer", "hoja"),
        [
            (leer_choferes, HOJA_CHOFERES),
            (leer_clientes, HOJA_CLIENTES),
            (leer_viajes, HOJA_VIAJES),
        ],
    )
    def test_cada_hoja_devuelve_su_numero_de_filas(self, leer, hoja: str, archivo: Path) -> None:
        filas = leer(archivo)

        assert len(filas) == ESPERADAS[hoja]

    def test_los_datos_vienen_por_nombre_de_columna(self, archivo: Path) -> None:
        """FR-002: si mañana cambia el orden de las columnas, el script sigue funcionando."""
        filas = leer_clientes(archivo)

        assert filas[0].datos["ID_Cliente"] == "CLI-101"
        assert filas[0].datos["Nombre_o_Empresa"] == "Muebles Belgrano S.R.L."

    def test_cada_fila_conserva_su_numero_original(self, archivo: Path) -> None:
        """FR-002 / FR-035: el reporte cita el número de fila del Excel, no uno del script."""
        filas = leer_clientes(archivo)

        assert filas[0].numero == 2
        assert filas[-1].numero == 11
        assert filas[0].hoja == HOJA_CLIENTES

    def test_los_renglones_vacios_no_se_cuentan(self, tmp_path: Path) -> None:
        """Un renglón en blanco al final de la hoja no es un cliente sin nombre."""
        from openpyxl import Workbook

        libro = Workbook()
        hoja = libro.active
        hoja.title = HOJA_CLIENTES
        hoja.append(["ID_Cliente", "Nombre_o_Empresa"])
        hoja.append(["CLI-101", "Uno"])
        hoja.append([None, None])
        hoja.append(["   ", "  "])
        ruta = tmp_path / "con_vacios.xlsx"
        libro.save(ruta)

        assert len(leer_clientes(ruta)) == 1

    def test_una_hoja_sin_filas_de_datos_falla(self, tmp_path: Path) -> None:
        from openpyxl import Workbook

        libro = Workbook()
        hoja = libro.active
        hoja.title = HOJA_CLIENTES
        hoja.append(["ID_Cliente", "Nombre_o_Empresa"])
        ruta = tmp_path / "solo_encabezado.xlsx"
        libro.save(ruta)

        with pytest.raises(ArchivoDeMigracionError, match="no tiene filas de datos"):
            leer_clientes(ruta)

    def test_el_valor_de_una_celda_se_conserva_como_esta(self, archivo: Path) -> None:
        """La extracción no normaliza: normalizar es otra etapa, con sus propias reglas."""
        filas = leer_choferes(archivo)

        assert filas[0].datos["Tipo_Vehiculo"] == "Furgón Grande"
        assert filas[0].datos["Tarifa_Hora_Base"] == "$12.500"
        assert filas[0].datos["Capacidad_Carga"] == "1500 kg"
